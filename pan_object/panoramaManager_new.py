#!/usr/bin/env python3
"""
Manage security rules and their referenced objects in Panorama device groups
using the pan-os-python SDK.

Configuration can be supplied either as an Excel workbook (.xlsx/.xlsm) or as
a JSON file (.json).

Excel format
------------
* Optional sheet named "settings" (or "config"/"meta"): rows of key/value
  pairs in columns A/B. Supported keys: ``audit_comment``.
* One sheet per object type / rulebase. Sheet names are matched
  case-insensitively (e.g. "address_object", "Address Objects", "pre-rulebase").
  Recognized types:
      address_object, address_group, service_object, service_group,
      url_category, external_dynamic_list,
      pre-rulebase, post-rulebase
* Each sheet must have a header row with a ``device_group`` column and a
  ``name`` column; any other columns map to pan-os-python attribute names
  (case-insensitive).
* List‑valued attributes (rule source/destination/service/etc., address group
  static_value, service group / url category value) are written as comma‑
  separated strings and automatically split into Python lists.
* Boolean attributes (disabled, negate_source, negate_destination) accept
  true/false, yes/no, 1/0.
* Rule ``move`` parameters are provided as a JSON string in a ``move`` column,
  e.g. ``{"where": "before", "destination": "some-rule"}``.
"""

import argparse
import getpass
import json
import logging
import os
import re
import sys
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional

from panos.objects import (
    AddressGroup,
    AddressObject,
    CustomUrlCategory,
    Edl,
    ServiceGroup,
    ServiceObject,
)
from panos.panorama import DeviceGroup, Panorama
from panos.policies import PostRulebase, PreRulebase, RuleAuditComment, SecurityRule

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
)
logger = logging.getLogger(__name__)

AUDIT_COMMENT_RE = re.compile(r"^(CHG|RITM|INC)[0-9]{7}")

# Registry of supported object types -> pan-os-python class
OBJECT_TYPES = {
    "address_object": AddressObject,
    "address_group": AddressGroup,
    "service_object": ServiceObject,
    "service_group": ServiceGroup,
    "url_category": CustomUrlCategory,
    "external_dynamic_list": Edl,
}

# Registry of supported rulebases -> pan-os-python class
RULEBASE_TYPES = {
    "pre-rulebase": PreRulebase,
    "post-rulebase": PostRulebase,
}

# Fields that should be treated as comma-separated lists, per config key.
LIST_FIELDS = {
    "address_group": {"static_value", "member_names"},
    "service_group": {"value"},
    "url_category": {"value"},
    "pre-rulebase": {
        "source",
        "destination",
        "source_hip",
        "destination_hip",
        "application",
        "service",
        "category",
        "source_user",
        "hip_profiles",
        "fromzone",
        "tozone",
        "tag",
    },
    "post-rulebase": {
        "source",
        "destination",
        "source_hip",
        "destination_hip",
        "application",
        "service",
        "category",
        "source_user",
        "hip_profiles",
        "fromzone",
        "tozone",
        "tag",
    },
}

BOOLEAN_FIELDS = {"disabled", "negate_source", "negate_destination"}

# Aliases for sheet names -> internal config keys.
SHEET_NAME_ALIASES = {
    "address_object": "address_object",
    "address_objects": "address_object",
    "addressobject": "address_object",
    "addresses": "address_object",
    "address_group": "address_group",
    "address_groups": "address_group",
    "addressgroup": "address_group",
    "service_object": "service_object",
    "service_objects": "service_object",
    "serviceobject": "service_object",
    "services": "service_object",
    "service_group": "service_group",
    "service_groups": "service_group",
    "servicegroup": "service_group",
    "url_category": "url_category",
    "url_categories": "url_category",
    "custom_url_category": "url_category",
    "custom_url_categories": "url_category",
    "urlcategory": "url_category",
    "external_dynamic_list": "external_dynamic_list",
    "external_dynamic_lists": "external_dynamic_list",
    "externaldynamiclist": "external_dynamic_list",
    "edl": "external_dynamic_list",
    "edls": "external_dynamic_list",
    "pre_rulebase": "pre-rulebase",
    "pre_rules": "pre-rulebase",
    "prerulebase": "pre-rulebase",
    "pre": "pre-rulebase",
    "post_rulebase": "post-rulebase",
    "post_rules": "post-rulebase",
    "postrulebase": "post-rulebase",
    "post": "post-rulebase",
}

SETTINGS_SHEET_NAMES = {"settings", "config", "meta", "metadata"}


def _log_section(message: str) -> None:
    logger.info("=" * 60)
    logger.info(message)
    logger.info("=" * 60)


def _normalize_sheet_name(name: str) -> Optional[str]:
    key = name.strip().lower().replace("-", "_").replace(" ", "_")
    key = re.sub(r"_+", "_", key)
    return SHEET_NAME_ALIASES.get(key)


def _clean_cell(value: Any) -> Any:
    """Normalize an Excel cell value: strip strings, stringify numbers."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, str):
        s = value.strip()
        return s if s else None
    return value


def _coerce_record_fields(object_key: str, record: Dict[str, Any]) -> Dict[str, Any]:
    """Convert string cell values to the types pan-os-python expects."""
    result = dict(record)

    for field in LIST_FIELDS.get(object_key, set()):
        if field in result and isinstance(result[field], str):
            result[field] = [
                item.strip() for item in result[field].split(",") if item.strip()
            ]

    for field in BOOLEAN_FIELDS:
        if field in result and isinstance(result[field], str):
            result[field] = result[field].strip().lower() in ("true", "yes", "1")

    if "move" in result and isinstance(result["move"], str):
        try:
            result["move"] = json.loads(result["move"])
        except json.JSONDecodeError:
            logger.warning("Invalid JSON in 'move' column: %r", result["move"])
            result.pop("move")

    return result


def _load_config_from_excel(path: Path) -> Dict[str, Any]:
    """Parse an Excel workbook into the cfg_data structure used by PanoramaManager."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "openpyxl is required to read Excel configs. Install with "
            "`pip install openpyxl`."
        ) from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    cfg_data: Dict[str, Any] = {}
    audit_comment: Optional[str] = None

    try:
        # ---- settings sheet (optional) --------------------------------
        for sheet_name in wb.sheetnames:
            if sheet_name.strip().lower() in SETTINGS_SHEET_NAMES:
                ws = wb[sheet_name]
                for row in ws.iter_rows(values_only=True):
                    if not row or row[0] is None:
                        continue
                    key = str(row[0]).strip().lower()
                    val = _clean_cell(row[1]) if len(row) > 1 else None
                    if key == "audit_comment" and val:
                        audit_comment = str(val)
                break

        # ---- object / rulebase sheets ---------------------------------
        for sheet_name in wb.sheetnames:
            if sheet_name.strip().lower() in SETTINGS_SHEET_NAMES:
                continue

            object_key = _normalize_sheet_name(sheet_name)
            if object_key is None:
                logger.warning("Skipping unrecognized sheet: %s", sheet_name)
                continue

            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                continue

            headers = [
                str(h).strip().lower() if h is not None else ""
                for h in header_row
            ]
            header_idx = {h: i for i, h in enumerate(headers) if h}

            if "device_group" not in header_idx:
                logger.warning(
                    "Sheet '%s' is missing a 'device_group' column, skipping",
                    sheet_name,
                )
                continue

            if "name" not in header_idx:
                logger.warning(
                    "Sheet '%s' is missing a 'name' column, skipping", sheet_name
                )
                continue

            for row in rows_iter:
                if not any(v is not None for v in row):
                    continue

                record: Dict[str, Any] = {}
                for h, idx in header_idx.items():
                    if idx >= len(row):
                        continue
                    cleaned = _clean_cell(row[idx])
                    if cleaned is None or cleaned == "":
                        continue
                    record[h] = cleaned

                device_group = record.pop("device_group", None)
                if not device_group:
                    logger.warning(
                        "Row without device_group on sheet '%s', skipping",
                        sheet_name,
                    )
                    continue

                if not record.get("name"):
                    logger.warning(
                        "Row without a name on sheet '%s', skipping", sheet_name
                    )
                    continue

                device_group = str(device_group).strip()
                record = _coerce_record_fields(object_key, record)

                cfg_data.setdefault(device_group, {}).setdefault(
                    object_key, []
                ).append(record)
    finally:
        wb.close()

    if audit_comment:
        cfg_data["audit_comment"] = audit_comment

    return cfg_data


def _load_config_from_json(path: Path) -> Dict[str, Any]:
    with open(path, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def _load_config(path: Path) -> Dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return _load_config_from_excel(path)
    if suffix == ".json":
        return _load_config_from_json(path)
    raise ValueError(
        f"Unsupported configuration file '{path.name}'. "
        "Expected .xlsx/.xlsm (Excel) or .json."
    )


class OperationType(Enum):
    CREATE = "create"
    DELETE = "delete"
    LIST = "list"
    MOVE = "move"

    @classmethod
    def from_string(cls, value: str) -> "OperationType":
        try:
            return cls(value.lower())
        except ValueError:
            raise ValueError(
                f"Invalid operation: {value}. "
                "Must be 'create', 'delete', 'list', or 'move'"
            )


class ObjectType(Enum):
    ADDRESS = "address_object"
    ADDRESS_GROUP = "address_group"
    SERVICE = "service_object"
    SERVICE_GROUP = "service_group"
    URL_CATEGORY = "url_category"
    EDL = "external_dynamic_list"

    @classmethod
    def from_string(cls, value: str) -> "ObjectType":
        try:
            return cls(value.lower())
        except ValueError:
            raise ValueError(f"Invalid object type: {value}.")


class RuleType(Enum):
    PRE_RULE = "pre-rulebase"
    POST_RULE = "post-rulebase"

    @classmethod
    def from_string(cls, value: str) -> "RuleType":
        try:
            return cls(value.lower())
        except ValueError:
            raise ValueError(
                f"Invalid rulebase type: {value}. "
                "Must be 'pre-rulebase' or 'post-rulebase'"
            )


class PanoramaManager:
    """Manage security rules and objects in a Panorama rulebase."""

    def __init__(
        self,
        hostname: str,
        username: Optional[str] = None,
        password: Optional[str] = None,
        api_key: Optional[str] = None,
        audit_comment: Optional[str] = None,
        commit_changes: bool = False,
        **kwargs,
    ):
        self.hostname = hostname
        self.username = username
        self.audit_comment = audit_comment
        self.commit_changes = commit_changes
        self.scope = None
        self.object_type = None
        self.rulebase = None

        if api_key:
            self.panorama = Panorama(hostname, api_key=api_key)
        elif username and password:
            self.panorama = Panorama(
                hostname, api_username=username, api_password=password
            )
        else:
            raise ValueError("Either api_key or username/password must be provided")

    # --- Lookup Scope ------------------------------------------------------------
    def _get_device_group(self, device_group_name: str) -> Optional[DeviceGroup]:
        for dg in DeviceGroup.refreshall(self.panorama):
            if dg.name == device_group_name:
                return dg
        return None

    def _get_existing_object(self, object_type: type, name: str):
        for obj in object_type.refreshall(self.scope):
            if obj.name == name:
                return obj
        return None

    def _get_existing_rule(self, rule_name: str) -> Optional[SecurityRule]:
        self.scope.add(self.rulebase)
        for rule in SecurityRule.refreshall(self.rulebase):
            if rule.name == rule_name:
                return rule
        return None

    def _set_scope(self, device_group_name: str) -> bool:
        if device_group_name == "Shared":
            self.scope = self.panorama
            return True
        device_group = self._get_device_group(device_group_name)
        if not device_group:
            logger.error(f"Error: '{device_group_name}' does not exist")
            return False
        self.scope = device_group
        return True

    def _upsert(self, object_type: type, name: str, params: Dict) -> bool:
        try:
            existing = self._get_existing_object(object_type, name)
            if existing:
                logger.info(
                    f"{object_type.__name__} '{name}' already exists, updating..."
                )
                for key, value in params.items():
                    if hasattr(existing, key):
                        setattr(existing, key, value)
                existing.apply()
                return bool(existing)

            logger.info(f"{object_type.__name__} '{name}' does not exist, creating...")
            new_obj = object_type(name=name, **params)
            self.scope.add(new_obj)
            new_obj.create()
            return bool(new_obj)
        except Exception as e:
            logger.error(f"Error managing {object_type.__name__} '{name}': {e}")
            return False

    def _delete_object(self, object_type: type, name: str) -> bool:
        try:
            existing = self._get_existing_object(object_type, name)
            if existing:
                existing.delete()
                logger.info(f"Deleted {object_type.__name__} '{name}'")
                return True
            logger.warning(f"{object_type.__name__} '{name}' not found")
            return False
        except Exception as e:
            logger.error(f"Error deleting {object_type.__name__} '{name}': {e}")
            return False

    @staticmethod
    def _normalize_object_params(object_key: str, params: Dict) -> Dict:
        """Map config-file field names to pan-os-python attribute names."""
        if object_key == "address_group":
            params = dict(params)
            if "filter_criteria" in params:
                params["dynamic_value"] = params.pop("filter_criteria")
            elif "member_names" in params:
                params["static_value"] = params.pop("member_names")
        return params

    def create_or_update_url_category(self, name, params):
        return self._upsert(CustomUrlCategory, name, params)

    def delete_url_category(self, name):
        return self._delete_object(CustomUrlCategory, name)

    def create_or_update_address_object(self, name, params):
        return self._upsert(AddressObject, name, params)

    def delete_address_object(self, name):
        return self._delete_object(AddressObject, name)

    def create_or_update_address_group(self, name, params):
        return self._upsert(
            AddressGroup, name, self._normalize_object_params("address_group", params)
        )

    def delete_address_group(self, name):
        return self._delete_object(AddressGroup, name)

    def create_or_update_service_object(self, name, params):
        return self._upsert(ServiceObject, name, params)

    def delete_service_object(self, name):
        return self._delete_object(ServiceObject, name)

    def create_or_update_service_group(self, name, params):
        return self._upsert(ServiceGroup, name, params)

    def delete_service_group(self, name):
        return self._delete_object(ServiceGroup, name)

    def create_or_update_edl(self, name, params):
        return self._upsert(Edl, name, params)

    def delete_edl(self, name):
        return self._delete_object(Edl, name)

    def create_or_update_rule(self, rule_params: Dict) -> bool:
        rule_name = rule_params.get("name")
        existing_rule = self._get_existing_rule(rule_name)

        try:
            if existing_rule:
                logger.info(f"Rule '{rule_name}' already exists, updating...")
                for key, value in rule_params.items():
                    if key != "name" and hasattr(existing_rule, key):
                        setattr(existing_rule, key, value)
                existing_rule.apply()
                RuleAuditComment(existing_rule).update(self.audit_comment)
                logger.info(f"Rule '{rule_name}' updated successfully.")
                return True

            logger.info(f"Rule '{rule_name}' does not exist, creating...")
            self.scope.add(self.rulebase)
            new_rule = SecurityRule(**rule_params)
            self.rulebase.add(new_rule)
            new_rule.create()
            RuleAuditComment(new_rule).update(self.audit_comment)
            logger.info(f"Rule '{rule_name}' created successfully.")
            return True
        except Exception as e:
            logger.error(f"Failed to upsert rule '{rule_name}': {e}")
            return False

    def move_rule(self, rule_name: str, move_params: Dict) -> bool:
        existing_rule = self._get_existing_rule(rule_name)
        if not existing_rule:
            logger.error(f"Rule '{rule_name}' does not exist.")
            return False
        try:
            existing_rule.move(**move_params)
            logger.info(f"Rule '{rule_name}' moved successfully.")
            return True
        except Exception as e:
            logger.error(f"Failed to move rule '{rule_name}': {e}")
            return False

    def delete_rule(self, rule_name: str) -> bool:
        existing_rule = self._get_existing_rule(rule_name)
        if not existing_rule:
            logger.error(f"Rule '{rule_name}' does not exist.")
            return False
        try:
            existing_rule.delete()
            logger.info(f"Rule '{rule_name}' deleted successfully.")
            return True
        except Exception as e:
            logger.error(f"Failed to delete rule '{rule_name}': {e}")
            return False

    def list_rule(self, rule_name: str) -> bool:
        existing_rule = self._get_existing_rule(rule_name)
        if existing_rule:
            logger.info(existing_rule.about())
            return True
        return False

    def _validate_audit_comment(self) -> bool:
        if not self.audit_comment:
            logger.error("Missing rule audit comment")
            return False
        if not AUDIT_COMMENT_RE.fullmatch(self.audit_comment):
            logger.error("Invalid rule audit comment")
            return False
        return True

    def _run_object_ops(
        self,
        operation: str,
        object_data: Dict,
        device_group_name: str,
        results: Dict[str, bool],
    ) -> None:
        for object_key, object_cls in OBJECT_TYPES.items():
            if object_key not in object_data:
                continue

            _log_section(
                f"{operation.capitalize()}ing {object_key}s in '{device_group_name}'"
            )
            self.object_type = object_cls

            for obj in object_data[object_key]:
                name = obj.get("name")
                if not name:
                    continue

                if operation == "create":
                    params = self._normalize_object_params(
                        object_key, {k: v for k, v in obj.items() if k != "name"}
                    )
                    success = self._upsert(object_cls, name, params)
                elif operation == "delete":
                    success = self._delete_object(object_cls, name)
                elif operation == "list":
                    existing = self._get_existing_object(object_cls, name)
                    success = existing is not None
                    if existing:
                        logger.info(existing.about())
                else:
                    continue

                results[f"{object_key}_{name}"] = success

    def _run_rule_ops(
        self,
        operation: str,
        object_data: Dict,
        device_group_name: str,
        results: Dict[str, bool],
    ) -> bool:
        """Returns False to signal early exit (e.g. bad audit comment)."""
        present = [(k, RULEBASE_TYPES[k]) for k in RULEBASE_TYPES if k in object_data]
        if not present:
            return True

        if operation == "create" and not self._validate_audit_comment():
            return False

        for rulebase_key, rulebase_cls in present:
            self.rulebase = rulebase_cls()
            _log_section(
                f"{operation.capitalize()}ing rules in '{device_group_name}'"
            )

            for rule in object_data[rulebase_key]:
                name = rule.get("name")
                if not name:
                    continue
                move_params = rule.get("move", {})

                if operation == "create":
                    rule_params = {k: v for k, v in rule.items() if k != "move" and v}
                    success = self.create_or_update_rule(rule_params)
                    if move_params:
                        success = self.move_rule(name, move_params)
                elif operation == "delete":
                    success = self.delete_rule(name)
                elif operation == "list":
                    success = self.list_rule(name)
                elif operation == "move":
                    if not move_params:
                        continue
                    success = self.move_rule(name, move_params)
                else:
                    continue

                results[f"{rulebase_key}_{name}"] = success

        return True

    def _maybe_commit(self, results: Dict[str, bool]) -> None:
        failures = [k for k, v in results.items() if v is False]
        if self.commit_changes and not failures:
            logger.info("Committing changes...")
            self.panorama.commit(admins=[self.username], sync=True)
            logger.info("Commit completed successfully")
        else:
            logger.info("Updated candidate configuration. Changes not committed")

    # --- Entry Point ---------------------------------------------------------------
    def run_operation(self, operation: str, cfg_data: Dict[str, Any]) -> Dict[str, bool]:
        results: Dict[str, bool] = {}

        try:
            op = OperationType.from_string(operation).value

            for device_group_name, object_data in cfg_data.items():
                if not self._set_scope(device_group_name):
                    return results
                if not object_data:
                    continue

                if op in ("create", "delete", "list"):
                    self._run_object_ops(op, object_data, device_group_name, results)

                if not self._run_rule_ops(
                    op, object_data, device_group_name, results
                ):
                    return results

            if op in ("create", "delete", "move") and results:
                self._maybe_commit(results)

            return results

        except Exception as e:
            logger.error(f"Error in object operation: {e}")
            return results


def parse_arguments():
    class Password(argparse.Action):
        def __call__(self, parser, namespace, values, option_string):
            if values is None:
                values = getpass.getpass()
            setattr(namespace, self.dest, values)

    parser = argparse.ArgumentParser(
        description="Arguments to run PanoramaManager script"
    )
    parser.add_argument(
        "--hostname", "-H", required=True,
        help="Panorama hostname or IP address",
    )
    parser.add_argument(
        "--username", "-u", type=str,
        help="Panorama admin username",
    )
    parser.add_argument(
        "--file", "-f", type=str,
        help="Configuration file name (.xlsx Excel workbook or .json)",
    )
    parser.add_argument(
        "--operation", "-o",
        choices=["create", "delete", "move", "list"],
        nargs="?", const="list", default="list",
        help="Operation modes are create, delete, move, or list. Default to 'list'",
    )

    auth = parser.add_mutually_exclusive_group(required=False)
    auth.add_argument(
        "--password", "-p", action=Password, nargs="?", dest="passwd",
        help="Panorama admin password",
    )
    auth.add_argument("--apikey", "-a", type=str, help="Panorama API key")

    parser.add_argument("--audit", type=str, help="Rule audit comments")
    parser.add_argument("--commit", action="store_true", help="Commit changes")

    return parser.parse_args()


def get_secret(vault, vaultpath):
    from encryption import CredentialManager
    return CredentialManager(vault, vaultpath).decrypt()


def main():
    args = parse_arguments()
    base_path = Path.home() / "pyenv3.13" / "panos" / "pan_project"
    filepath = Path(f"{base_path}/config/{args.file}")

    PANORAMA_HOST = args.hostname
    API_KEY = args.apikey
    USERNAME = args.username
    PASSWORD = args.passwd
    OPERATION = args.operation
    AUDIT_COMMENT = args.audit or None
    COMMIT = args.commit
    VAULT = "panos_secrets.bin"
    vaultpath = Path.home() / "pyenv3.13" / "secrets"

    if not os.path.isfile(filepath):
        logger.info("Error: Objects must be provided")
        sys.exit(0)

    try:
        data = _load_config(filepath)
    except Exception as e:
        logger.error(f"Failed to load configuration file '{filepath}': {e}")
        sys.exit(1)

    if not AUDIT_COMMENT:
        AUDIT_COMMENT = data.get("audit_comment")
    cfg_data = {k: v for k, v in data.items() if k != "audit_comment"}

    if not cfg_data:
        sys.exit(0)

    # Resolve credentials and construct the manager
    if not API_KEY and not USERNAME:
        API_KEY = get_secret(VAULT, vaultpath).get("pano_apikey")
    if USERNAME and not PASSWORD:
        PASSWORD = get_secret(VAULT, vaultpath).get(USERNAME)

    manager = PanoramaManager(
        hostname=PANORAMA_HOST,
        username=USERNAME,
        password=PASSWORD,
        api_key=API_KEY,
        audit_comment=AUDIT_COMMENT,
        commit_changes=COMMIT,
    )

    results = manager.run_operation(OPERATION, cfg_data)

    logger.info("Operation results:")
    for obj_name, success in results.items():
        logger.info(f"  {'✓' if success else '✗'} {obj_name}")

    _log_section("Operations completed successfully!")


if __name__ == "__main__":
    main()
